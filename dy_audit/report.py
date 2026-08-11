"""Per-loan findings workbook (spec section 0).

One `.xlsx` per loan, matching the source material so it can be read alongside
the model. A summary sheet up top, then every check with its severity, status,
cell reference and evidence, then the parameters read out of the loan agreement.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook as XlsxWorkbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import SEVERITY_ORDER, LoanResult, Severity, Status

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_TITLE_FONT = Font(bold=True, size=14)
_LABEL_FONT = Font(bold=True)
_WRAP = Alignment(vertical="top", wrap_text=True)
_TOP = Alignment(vertical="top")

#: Severity colours, loudest first.
_SEVERITY_FILL = {
    Severity.BLOCKER: PatternFill("solid", fgColor="C00000"),
    Severity.HIGH: PatternFill("solid", fgColor="E8A33D"),
    Severity.MEDIUM: PatternFill("solid", fgColor="FFD966"),
    Severity.LOW: PatternFill("solid", fgColor="D9E2F3"),
    Severity.STANDING: PatternFill("solid", fgColor="C6E0B4"),
    Severity.INFO: PatternFill("solid", fgColor="F2F2F2"),
}
_SEVERITY_FONT = {
    Severity.BLOCKER: Font(bold=True, color="FFFFFF"),
    Severity.HIGH: Font(bold=True),
}

#: Status colours. MANUAL_REVIEW must never look like a pass.
_STATUS_FILL = {
    Status.FLAG: PatternFill("solid", fgColor="F4B6B6"),
    Status.PASS: PatternFill("solid", fgColor="D7E8D4"),
    Status.MANUAL_REVIEW: PatternFill("solid", fgColor="FFE699"),
    Status.UNVERIFIABLE: PatternFill("solid", fgColor="D0CECE"),
}

_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_FINDING_COLUMNS = [
    ("Severity", 11),
    ("Status", 16),
    ("Check", 28),
    ("Sheet", 22),
    ("Cell", 9),
    ("On DY path", 11),
    ("Finding", 82),
    ("Evidence", 74),
]


def _sort_key(finding):
    # Flags first within each severity, so the reader meets the actions first.
    status_rank = {Status.FLAG: 0, Status.MANUAL_REVIEW: 1, Status.UNVERIFIABLE: 2, Status.PASS: 3}
    return (
        SEVERITY_ORDER.get(finding.severity, 9),
        status_rank.get(finding.status, 9),
        finding.check_id,
    )


def _write_row(sheet, row: int, label: str, value, number_format: str | None = None) -> int:
    sheet.cell(row=row, column=1, value=label).font = _LABEL_FONT
    cell = sheet.cell(row=row, column=2, value=value)
    if number_format:
        cell.number_format = number_format
    cell.alignment = _TOP
    return row + 1


def _summary_sheet(sheet, result: LoanResult, run_date: dt.date) -> None:
    facts = result.facts
    sheet.cell(row=1, column=1, value=f"Debt Yield Audit - {result.loan_name}").font = _TITLE_FONT

    row = 3
    row = _write_row(sheet, row, "Run date", run_date.strftime("%Y-%m-%d"))
    row = _write_row(sheet, row, "Workbook", facts.get("workbook"))
    row = _write_row(sheet, row, "Definitions", facts.get("definitions"))
    row = _write_row(sheet, row, "OSAR tab audited", facts.get("osar_tab"))
    row = _write_row(sheet, row, "DY-test column", facts.get("dy_column"))
    period = facts.get("period_end")
    row = _write_row(sheet, row, "Statement ending", period.strftime("%m/%d/%Y") if period else "not read")
    hidden = facts.get("hidden_osar_tabs") or []
    row = _write_row(
        sheet, row, "Hidden OSAR tabs (ignored)", ", ".join(hidden) if hidden else "none"
    )

    row += 1
    sheet.cell(row=row, column=1, value="Reported figures").font = _HEADER_FONT
    sheet.cell(row=row, column=1).fill = _HEADER_FILL
    sheet.cell(row=row, column=2).fill = _HEADER_FILL
    row += 1
    row = _write_row(sheet, row, "Debt yield (NCF / UPB)", facts.get("debt_yield"), "0.0000%")
    row = _write_row(sheet, row, "Net cash flow (NCF)", facts.get("ncf"), "#,##0.00")
    row = _write_row(sheet, row, "Net operating income", facts.get("noi"), "#,##0.00")
    row = _write_row(sheet, row, "Effective gross income", facts.get("egi"), "#,##0.00")
    row = _write_row(sheet, row, "Occupancy", facts.get("occupancy"), "0.00%")
    row = _write_row(sheet, row, "UPB used", facts.get("upb"), "#,##0.00")
    row = _write_row(sheet, row, "UPB cell", facts.get("upb_cell"))
    sheet.cell(row=row, column=1, value="UPB is echoed, not verified").font = _LABEL_FONT
    sheet.cell(
        row=row,
        column=2,
        value="There is no source file for unpaid principal balance - confirm against internal records.",
    ).alignment = _TOP
    row += 2

    sheet.cell(row=row, column=1, value="Findings by severity").font = _HEADER_FONT
    sheet.cell(row=row, column=1).fill = _HEADER_FILL
    sheet.cell(row=row, column=2).fill = _HEADER_FILL
    row += 1
    counts = result.count_by_severity()
    for severity in sorted(Severity, key=lambda s: SEVERITY_ORDER.get(s, 9)):
        count = counts.get(severity, 0)
        sheet.cell(row=row, column=1, value=severity.value).font = _LABEL_FONT
        if count:
            sheet.cell(row=row, column=1).fill = _SEVERITY_FILL[severity]
            font = _SEVERITY_FONT.get(severity)
            if font:
                sheet.cell(row=row, column=1).font = font
        sheet.cell(row=row, column=2, value=count)
        row += 1

    row += 1
    sources = facts.get("source_tabs") or {}
    sheet.cell(row=row, column=1, value="Supporting tabs used").font = _HEADER_FONT
    sheet.cell(row=row, column=1).fill = _HEADER_FILL
    sheet.cell(row=row, column=2).fill = _HEADER_FILL
    row += 1
    for role, label in (("rent_roll", "Rent roll"), ("t12", "T12 / operating statement"), ("ar", "AR aging")):
        row = _write_row(sheet, row, label, sources.get(role) or "not present in this workbook")

    if result.error:
        row += 1
        sheet.cell(row=row, column=1, value="PROCESSING ERROR").font = Font(bold=True, color="C00000")
        sheet.cell(row=row, column=2, value=result.error).alignment = _WRAP

    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 82
    sheet.freeze_panes = "A3"


def _findings_sheet(sheet, result: LoanResult) -> None:
    for index, (title, width) in enumerate(_FINDING_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _TOP
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row_index, finding in enumerate(sorted(result.findings, key=_sort_key), start=2):
        values = [
            finding.severity.value,
            finding.status.value,
            finding.check_id,
            finding.sheet or "",
            finding.cell or "",
            "" if finding.on_dy_path is None else ("yes" if finding.on_dy_path else "no"),
            finding.message,
            finding.evidence or "",
        ]
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.alignment = _WRAP if column_index >= 7 else _TOP
            cell.border = _BORDER
        severity_cell = sheet.cell(row=row_index, column=1)
        severity_cell.fill = _SEVERITY_FILL[finding.severity]
        font = _SEVERITY_FONT.get(finding.severity)
        if font:
            severity_cell.font = font
        sheet.cell(row=row_index, column=2).fill = _STATUS_FILL[finding.status]

    sheet.freeze_panes = "A2"
    if result.findings:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(_FINDING_COLUMNS))}{len(result.findings) + 1}"


def _parameters_sheet(sheet, result: LoanResult) -> None:
    params = result.facts.get("params")
    headers = ["Parameter", "Value", "Quoted from the loan agreement"]
    widths = [30, 26, 108]
    for index, (title, width) in enumerate(zip(headers, widths), start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        sheet.column_dimensions[get_column_letter(index)].width = width

    if params is None:
        sheet.cell(row=2, column=1, value="Definitions file was not parsed for this loan.")
        return

    rows = [
        ("Vacancy floor", params.vacancy_floor),
        ("Occupancy threshold", None),
        ("Investment-grade carve-out", params.ig_vacancy_carveout),
        ("Reserve rate", params.reserve_rate),
        ("Reserve basis", params.reserve_basis),
        ("Management fee %", params.mgmt_fee_pct),
        ("Base stated in agreement", params.mgmt_stated_base),
        ("Base actually enforced", None),
        ("Delinquency window", params.delinquency),
        ("Other income window (months)", params.other_income_months),
        ("Concessions window (months)", params.concession_months),
        ("New-lease occupancy window", params.new_lease_days),
        ("New-lease window, investment grade", params.new_lease_ig_days),
        ("Rent-step window", params.rent_step_window),
        ("Known-vacate window", params.known_vacate_days),
    ]
    row_index = 2
    for label, parsed in rows:
        sheet.cell(row=row_index, column=1, value=label).font = _LABEL_FONT
        if label == "Occupancy threshold":
            threshold = params.occupancy_threshold
            sheet.cell(row=row_index, column=2, value=f"{threshold:.2%}" if threshold else "n/a")
            sheet.cell(row=row_index, column=3, value="derived as 1 - vacancy floor").alignment = _WRAP
        elif label == "Base actually enforced":
            sheet.cell(row=row_index, column=2, value=params.mgmt_base)
            sheet.cell(
                row=row_index,
                column=3,
                value="House rule: the 3% base is always EGI regardless of the agreement's wording.",
            ).alignment = _WRAP
        else:
            sheet.cell(row=row_index, column=2, value=str(parsed))
            sheet.cell(row=row_index, column=3, value=parsed.quote or "").alignment = _WRAP
        row_index += 1

    if params.conflicts:
        row_index += 1
        sheet.cell(row=row_index, column=1, value="Conflicts with the build spec").font = _HEADER_FONT
        sheet.cell(row=row_index, column=1).fill = _HEADER_FILL
        row_index += 1
        for conflict in params.conflicts:
            sheet.cell(row=row_index, column=1, value="Loan agreement governs").font = _LABEL_FONT
            sheet.cell(row=row_index, column=2, value=conflict).alignment = _WRAP
            row_index += 1
    sheet.freeze_panes = "A2"


_REBUILD_COLUMNS = [
    ("Row", 7),
    ("Unit", 12),
    ("Tenant", 30),
    ("Status", 18),
    ("Monthly rent", 13),
    ("Annualized", 14),
    ("Included", 10),
    ("Why", 44),
    ("Cell is", 20),
    ("Balance", 11),
    ("Move-out", 11),
]

_INCLUDED_FILL = PatternFill("solid", fgColor="D7E8D4")
_EXCLUDED_FILL = PatternFill("solid", fgColor="F4B6B6")
_EDIT_FILL = PatternFill("solid", fgColor="FFE699")


def _rebuild_sheet(sheet, result: LoanResult) -> None:
    """The independent rent-roll rebuild: every row, every decision, and the
    reconciliation against the OSAR when the totals do not tie."""
    rebuilds = result.facts.get("rebuilds") or []
    if not rebuilds:
        sheet.cell(row=1, column=1, value="No rent roll could be rebuilt for this loan.")
        return

    row_index = 1
    for rebuild in rebuilds:
        sheet.cell(
            row=row_index,
            column=1,
            value=f"{rebuild.line.value} rebuilt from {rebuild.sheet!r}",
        ).font = _TITLE_FONT
        row_index += 1
        scale = rebuild.scale
        header_bits = [
            f"rent column {rebuild.columns.get('rent', '?')} ({rebuild.rent_header})",
            "monthly x 12" if rebuild.monthly else "annual",
            f"source: {rebuild.source}",
        ]
        if rebuild.as_of:
            header_bits.append(rebuild.as_of)
        sheet.cell(row=row_index, column=1, value="; ".join(header_bits)).alignment = _WRAP
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=8)
        row_index += 2

        head_row = row_index
        for index, (title, width) in enumerate(_REBUILD_COLUMNS, start=1):
            cell = sheet.cell(row=head_row, column=index, value=title)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            column = get_column_letter(index)
            if width > (sheet.column_dimensions[column].width or 0):
                sheet.column_dimensions[column].width = width
        row_index += 1

        for tenant in rebuild.rows:
            values = [
                tenant.row,
                tenant.unit,
                tenant.name,
                tenant.status,
                tenant.monthly_rent if rebuild.monthly else tenant.monthly_rent / 12.0,
                tenant.monthly_rent * scale,
                "yes" if tenant.included else "no",
                tenant.reason,
                "formula: " + (tenant.rent_formula or "")[:40] if tenant.rent_is_formula else "value",
                tenant.balance,
                tenant.move_out.strftime("%m/%d/%Y") if tenant.move_out else None,
            ]
            for column_index, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_index, column=column_index, value=value)
                cell.border = _BORDER
                if column_index in (5, 6, 10):
                    cell.number_format = "#,##0.00"
                if column_index == 8:
                    cell.alignment = _WRAP
            sheet.cell(row=row_index, column=7).fill = (
                _INCLUDED_FILL if tenant.included else _EXCLUDED_FILL
            )
            if tenant.rent_is_formula and rebuild.source == "headers":
                sheet.cell(row=row_index, column=9).fill = _EDIT_FILL
            row_index += 1

        row_index += 1
        rebuilt = rebuild.annualized()
        model = rebuild.model_gpr
        sheet.cell(row=row_index, column=1, value="Rebuilt annualized base rent").font = _LABEL_FONT
        cell = sheet.cell(row=row_index, column=6, value=rebuilt)
        cell.number_format = "#,##0.00"
        cell.font = _LABEL_FONT
        row_index += 1
        sheet.cell(row=row_index, column=1, value=f"Model {rebuild.line.value} line").font = _LABEL_FONT
        cell = sheet.cell(row=row_index, column=6, value=model)
        cell.number_format = "#,##0.00"
        row_index += 1
        if model is not None:
            difference = model - rebuilt
            ties = abs(difference) <= max(1.0, abs(model) * 0.001)
            label = "TIES within 0.1%" if ties else f"DOES NOT TIE ({difference:+,.0f})"
            cell = sheet.cell(row=row_index, column=1, value=label)
            cell.font = Font(bold=True, color="376E37" if ties else "C00000")
            row_index += 1
        if rebuild.model_note:
            row_index = _write_row(sheet, row_index, "OSAR note on this line", rebuild.model_note)

        if rebuild.reconciliation:
            row_index += 1
            cell = sheet.cell(row=row_index, column=1, value="Reconciliation of the difference")
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            sheet.cell(row=row_index, column=2).fill = _HEADER_FILL
            row_index += 1
            for item, amount in rebuild.reconciliation:
                sheet.cell(row=row_index, column=1, value=item).alignment = _WRAP
                sheet.merge_cells(
                    start_row=row_index, start_column=1, end_row=row_index, end_column=5
                )
                if amount is not None:
                    cell = sheet.cell(row=row_index, column=6, value=amount)
                    cell.number_format = "#,##0.00"
                row_index += 1

        excluded = rebuild.excluded_with_rent()
        if excluded:
            row_index += 1
            sheet.cell(
                row=row_index, column=1, value="Excluded rows that still carry rent"
            ).font = _LABEL_FONT
            row_index += 1
            for status, (count, total) in sorted(excluded.items()):
                sheet.cell(row=row_index, column=1, value=f"{status}: {count} row(s)")
                cell = sheet.cell(row=row_index, column=6, value=total * scale)
                cell.number_format = "#,##0.00"
                row_index += 1

        for note in rebuild.notes:
            sheet.cell(row=row_index, column=1, value=f"Note: {note}").alignment = _WRAP
            sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=8)
            row_index += 1
        row_index += 2

    sheet.freeze_panes = "A2"


#: Phase 0's term sheet, in the order the prompt lists it. The labels are the
#: reviewer's words rather than the schema's keys.
_TERM_SHEET_ROWS = [
    ("delinquency_threshold_days", "Delinquency threshold (days)"),
    ("notice_to_vacate", "Notice to vacate"),
    ("vacancy_floor_pct", "Vacancy floor"),
    ("vacancy_floor_base", "Vacancy floor applies to"),
    ("occupancy_cap_pct", "Occupancy cap"),
    ("new_lease_window_days", "New / signed lease window (days)"),
    ("concessions_basis", "Concessions basis"),
    ("other_income_basis", "Other income basis"),
    ("tenant_status_screens", "Tenant-status screens stated"),
]

_REVIEW_REBUILD_COLUMNS = [
    ("Revenue line", 28),
    ("OSAR cell", 12),
    ("Reported", 15),
    ("Rebuilt", 15),
    ("Variance", 14),
    ("Ties", 8),
    ("How it was rebuilt", 70),
]

_REVIEW_FINDING_COLUMNS = [
    ("ID", 6),
    ("Rule", 7),
    ("Severity", 10),
    ("Type", 13),
    ("Revenue line", 22),
    ("Cells", 26),
    ("Finding", 62),
    ("Impact (annual)", 15),
    ("Clause relied on", 46),
    ("Evidence", 46),
    ("Recommendation", 44),
]


def _section(sheet, row: int, title: str, width: int) -> int:
    """A header band matching the one the other sheets use."""
    cell = sheet.cell(row=row, column=1, value=title)
    cell.font = _HEADER_FONT
    cell.fill = _HEADER_FILL
    for column in range(2, width + 1):
        sheet.cell(row=row, column=column).fill = _HEADER_FILL
    return row + 1


def _table_head(sheet, row: int, columns) -> int:
    for index, (title, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=row, column=index, value=title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _WRAP
        sheet.column_dimensions[cell.column_letter].width = width
    return row + 1


def _revenue_sheet(sheet, result: LoanResult) -> None:
    """The half of the LLM review the eight-column Findings sheet cannot hold.

    The Findings sheet stays the report's spine and is unchanged. What lands
    here is everything the review returns that has no column there: the term
    sheet it built from the agreement, the line-by-line rebuild, the reviewer's
    prose note, and the full record behind each finding - rule number, clause,
    impact, recommendation, and every cell rather than just the first.
    """
    payload = result.facts.get("revenue_review")
    if not isinstance(payload, dict):
        sheet.cell(
            row=1,
            column=1,
            value=(
                "No model revenue review is attached to this loan. Either the run "
                "used --no-llm, in which case the deterministic revenue checks "
                "answered these lines and appear on the Findings sheet, or the "
                "review did not complete - look for CHK_REVENUE_LLM."
            ),
        ).alignment = _WRAP
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
        sheet.column_dimensions["A"].width = 110
        return

    row = 1
    sheet.cell(row=row, column=1, value="Revenue review").font = _TITLE_FONT
    row += 1
    header_bits = [
        f"loan: {payload.get('loan') or result.loan_name}",
        f"test date: {payload.get('test_date') or 'not stated'}",
        f"OSAR tab: {payload.get('osar_tab') or 'not stated'}",
    ]
    sheet.cell(row=row, column=1, value="; ".join(header_bits)).alignment = _WRAP
    row += 2

    note = (payload.get("reviewer_note") or "").strip()
    if note:
        row = _section(sheet, row, "Reviewer note", 7)
        cell = sheet.cell(row=row, column=1, value=note)
        cell.alignment = _WRAP
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        sheet.row_dimensions[row].height = 108
        row += 2

    term_sheet = payload.get("term_sheet")
    if isinstance(term_sheet, dict):
        row = _section(sheet, row, "Term sheet, as read from the agreement", 7)
        sheet.cell(
            row=row,
            column=1,
            value=(
                "The model's reading of the clause governs. A blank means the "
                "agreement states no such parameter - which is a real answer, and "
                "stands the rules that depend on it down. Compare against the "
                "parsed values on the Loan Parameters sheet."
            ),
        ).alignment = _WRAP
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        row += 1
        for key, label in _TERM_SHEET_ROWS:
            value = term_sheet.get(key)
            rendered = "not stated in the agreement" if value is None else value
            sheet.cell(row=row, column=1, value=label).font = _LABEL_FONT
            cell = sheet.cell(row=row, column=2, value=rendered)
            cell.alignment = _WRAP
            if value is None:
                cell.font = Font(italic=True, color="808080")
            elif key.endswith("_pct") and isinstance(value, (int, float)):
                cell.number_format = "0.00%"
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
            row += 1
        row += 1

    rebuild = payload.get("rebuild")
    if isinstance(rebuild, list) and rebuild:
        row = _section(sheet, row, "Independent rebuild of each revenue line", 7)
        row = _table_head(sheet, row, _REVIEW_REBUILD_COLUMNS)
        for entry in rebuild:
            if not isinstance(entry, dict):
                continue
            ties = str(entry.get("flag") or "").upper() == "PASS"
            values = [
                entry.get("line"),
                entry.get("osar_cell"),
                entry.get("reported"),
                entry.get("rebuilt"),
                entry.get("variance"),
                "PASS" if ties else "FLAG",
                entry.get("derivation"),
            ]
            for index, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=index, value=value)
                cell.alignment = _WRAP if index == 7 else _TOP
                cell.border = _BORDER
                if index in (3, 4, 5) and isinstance(value, (int, float)):
                    cell.number_format = "#,##0.00"
            flag_cell = sheet.cell(row=row, column=6)
            flag_cell.fill = _STATUS_FILL[Status.PASS if ties else Status.FLAG]
            flag_cell.font = Font(bold=True, color="376E37" if ties else "C00000")
            row += 1
        row += 2

    findings = payload.get("findings")
    if isinstance(findings, list) and findings:
        row = _section(sheet, row, "Findings in full", 11)
        head_row = row
        row = _table_head(sheet, row, _REVIEW_FINDING_COLUMNS)
        for entry in findings:
            if not isinstance(entry, dict):
                continue
            cells = entry.get("cells")
            if isinstance(cells, str):
                cells = [cells]
            values = [
                entry.get("id"),
                entry.get("rule"),
                entry.get("severity"),
                entry.get("type"),
                entry.get("revenue_line"),
                ", ".join(c for c in (cells or []) if isinstance(c, str)),
                entry.get("finding"),
                entry.get("impact_annualized"),
                entry.get("clause"),
                entry.get("evidence"),
                entry.get("recommendation"),
            ]
            for index, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=index, value=value)
                cell.alignment = _WRAP
                cell.border = _BORDER
                if index == 8 and isinstance(value, (int, float)):
                    cell.number_format = "#,##0.00"
            severity = _coerce_severity(entry.get("severity"))
            if severity is not None:
                severity_cell = sheet.cell(row=row, column=3)
                severity_cell.fill = _SEVERITY_FILL[severity]
                if severity in _SEVERITY_FONT:
                    severity_cell.font = _SEVERITY_FONT[severity]
            row += 1
        sheet.freeze_panes = sheet.cell(row=head_row + 1, column=1).coordinate

    sheet.column_dimensions["A"].width = max(sheet.column_dimensions["A"].width or 0, 28)


def _coerce_severity(value) -> Severity | None:
    """The review writes 'High'; the palette is keyed on the Severity enum."""
    try:
        return Severity(str(value).strip().upper())
    except (ValueError, AttributeError):
        return None


def write_findings_workbook(result: LoanResult, destination: Path, run_date: dt.date | None = None) -> Path:
    """Write one loan's findings workbook and return the path written."""
    run_date = run_date or dt.date.today()
    workbook = XlsxWorkbook()

    summary = workbook.active
    summary.title = "Summary"
    _summary_sheet(summary, result, run_date)
    _findings_sheet(workbook.create_sheet("Findings"), result)
    _rebuild_sheet(workbook.create_sheet("Rent Roll Rebuild"), result)
    _revenue_sheet(workbook.create_sheet("Revenue Review"), result)
    _parameters_sheet(workbook.create_sheet("Loan Parameters"), result)

    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination


def findings_filename(loan_name: str, run_date: dt.date) -> str:
    safe = "".join(c if c.isalnum() or c in " -_" else "_" for c in loan_name).strip()
    return f"{safe} - DY Audit Findings - {run_date:%Y_%m_%d}.xlsx"

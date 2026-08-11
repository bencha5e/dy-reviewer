"""Independent recompute of the two lines the spec requires rebuilt: GPR and vacancy.

Everything else in the tool verifies the model's own formulas. These two are
rebuilt from the rent roll and compared back, because they are where the
definitional rules actually bite.

The governing rule is **never guess**. A rent roll the parser cannot read
confidently produces MANUAL_REVIEW, never a variance flag: a half-parsed rent
roll reporting a spurious GPR difference would be worse than no check at all.
Rent steps and free rent are deliberately left to the reviewer (spec section 5,
rule 12), so a model carrying those adjustments is reported with both figures
rather than flagged.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl.utils import column_index_from_string, get_column_letter

from . import formula as F
from .context import LoanContext
from .model import Finding, Severity, Status
from .osar import Line

#: Individual revenue lines flag above 0.1% difference (spec section 4).
REVENUE_TOLERANCE = 0.001

#: A status that means the tenant is in occupancy and paying.
_OCCUPIED = re.compile(r"^occupied", re.I)
#: Columns whose header marks an analyst include/exclude decision.
_INCLUDE_HEADER = re.compile(r"include.*(dy|test)|(dy|test).*include", re.I)
#: Headers signalling that the model applies step-up or free-rent adjustments,
#: which the definition permits and the tool does not attempt to verify.
_ADJUSTMENT_HEADER = re.compile(
    r"free rent|step[- ]?up|months? included|months? excluded|adj\w*\s+annual", re.I
)
_MONTHLY_HEADER = re.compile(r"month", re.I)
_ANNUAL_HEADER = re.compile(r"annual|annualiz", re.I)

_CELL_RE = re.compile(r"^([A-Z]+)(\d+)$")


def _split(coord: str) -> tuple[str, int]:
    m = _CELL_RE.match(coord.replace("$", "").upper())
    return (m.group(1), int(m.group(2))) if m else ("A", 0)


@dataclass
class RentRollParse:
    """What the parser could establish about a rent roll."""

    sheet: str | None = None
    rent_column: str | None = None
    rows: list[int] = field(default_factory=list)
    values: dict[int, float] = field(default_factory=dict)
    status_column: str | None = None
    statuses: dict[int, str] = field(default_factory=dict)
    monthly: bool = True
    has_adjustments: bool = False
    adjustment_headers: list[str] = field(default_factory=list)
    #: True when the summed column is itself computed by the model (an "adjusted
    #: rent" column) rather than raw contractual rent. The recompute then
    #: re-sums the model's own working rather than rebuilding from scratch.
    rent_column_is_derived: bool = False
    notes: list[str] = field(default_factory=list)

    @property
    def confident(self) -> bool:
        return bool(self.sheet and self.rent_column and self.values)

    def included_rows(self) -> list[int]:
        """Rows counted as in-place, current tenants."""
        if not self.statuses:
            return sorted(self.values)
        return sorted(r for r in self.values if _OCCUPIED.match(self.statuses.get(r, "")))

    def annualized(self) -> float:
        total = sum(self.values[r] for r in self.included_rows())
        return total * 12 if self.monthly else total


def _header_text(ctx: LoanContext, sheet: str, column: str, first_row: int) -> str:
    """Join the header cells sitting above a column's data."""
    parts = []
    for row in range(max(1, first_row - 6), first_row):
        text = ctx.wb.text(sheet, f"{column}{row}")
        if text:
            parts.append(text)
    return " ".join(parts)


def _candidate_ranges(ctx: LoanContext, sheet: str, start_sheet: str, formula: str | None,
                      depth: int = 0) -> list[F.Ref]:
    """Ranges on the rent-roll sheet reachable from a formula.

    Strada's GPR reaches its tenant rows through a summary block of `SUMIF`s, so
    the search follows single-cell references on the rent roll one more hop.
    """
    if not formula or depth > 3:
        return []
    found: list[F.Ref] = []
    for ref in F.iter_refs(formula):
        ref_sheet = ref.resolved_sheet(start_sheet)
        if ref_sheet != sheet:
            continue
        if ref.is_range:
            found.append(ref)
        else:
            found.extend(
                _candidate_ranges(ctx, sheet, sheet, ctx.wb.formula(sheet, ref.coord), depth + 1)
            )
    return found


def parse_rent_roll(ctx: LoanContext) -> RentRollParse:
    """Locate the tenant rows and the in-place rent column feeding GPR."""
    wb, tab = ctx.wb, ctx.tab
    parse = RentRollParse()
    tabs = ctx.facts.get("source_tabs") or {}
    sheet = tabs.get("rent_roll")
    if not sheet or not wb.has_sheet(sheet):
        parse.notes.append("could not identify the rent-roll tab from the GPR formula")
        return parse
    parse.sheet = sheet

    gpr_coord = tab.cell(Line.GPR)
    if gpr_coord is None:
        parse.notes.append("no GPR line on the OSAR tab")
        return parse
    res = F.resolve_defining_formula(wb, tab.sheet, gpr_coord)

    ranges = _candidate_ranges(ctx, sheet, res.sheet, res.formula)
    # A pass-through that lands directly on a rent-roll cell (Ares: SUM(RR!J8))
    # leaves no range; treat the landing cell's column as the rent column.
    if not ranges and res.sheet == sheet:
        col, row = _split(res.coord)
        parse.rent_column, parse.rows = col, [row]
        value = wb.number(sheet, res.coord)
        if value is not None:
            parse.values[row] = value

    numeric_ranges: list[tuple[F.Ref, dict[int, float]]] = []
    text_ranges: list[tuple[F.Ref, dict[int, str]]] = []
    rows_bound, _cols = wb.bounds(sheet)

    for ref in ranges:
        body = ref.body.replace("$", "").upper()
        start, end = body.split(":")
        if ref.is_whole_column:
            col = re.sub(r"\d", "", start)
            row_span = range(1, min(rows_bound, 2000) + 1)
        else:
            c1, r1 = _split(start)
            c2, r2 = _split(end)
            if c1 != c2:
                continue
            col, row_span = c1, range(min(r1, r2), max(r1, r2) + 1)
        numbers: dict[int, float] = {}
        texts: dict[int, str] = {}
        for row in row_span:
            value = wb.value(sheet, f"{col}{row}")
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                numbers[row] = float(value)
            elif isinstance(value, str) and value.strip():
                texts[row] = value.strip()
        if numbers:
            numeric_ranges.append((ref, numbers))
        if texts and len(texts) > len(numbers):
            text_ranges.append((ref, texts))

    # Resolve the status column before the rent column: on a unit-level rent roll
    # the two must describe the same rows. Strada reaches its tenant rents
    # through a summary block of SUMIFs, so both the tenant range (T7:T517) and
    # the summary range (M531:M539) are reachable and sum to the same total -
    # only the tenant range lines up with the statuses in H7:H517.
    for ref, texts in text_ranges:
        if any(_OCCUPIED.match(s) for s in list(texts.values())[:60]):
            parse.status_column = re.sub(
                r"\d", "", ref.body.replace("$", "").upper().split(":")[0]
            )
            parse.statuses = texts
            break

    if numeric_ranges and not parse.rent_column:
        if parse.statuses:
            status_rows = set(parse.statuses)
            aligned = [
                (ref, numbers) for ref, numbers in numeric_ranges if status_rows & set(numbers)
            ]
            numeric_ranges = aligned or numeric_ranges

        # Pick the numeric range that best explains the model's GPR, testing both
        # a monthly and an already-annual reading. Scoring uses the filtered
        # total so the definitional inclusion rule is part of the match.
        model_gpr = wb.number(tab.sheet, gpr_coord) or 0.0
        best = None
        for ref, numbers in numeric_ranges:
            if parse.statuses:
                rows = [r for r in numbers if _OCCUPIED.match(parse.statuses.get(r, ""))]
                total = sum(numbers[r] for r in rows) if rows else sum(numbers.values())
            else:
                total = sum(numbers.values())
            for monthly, scaled in ((True, total * 12), (False, total)):
                if scaled == 0:
                    continue
                error = abs(scaled - model_gpr) / max(abs(model_gpr), 1.0)
                if best is None or error < best[0]:
                    best = (error, ref, numbers, monthly)
        if best:
            _err, ref, numbers, monthly = best
            parse.rent_column = re.sub(r"\d", "", ref.body.replace("$", "").upper().split(":")[0])
            parse.values = numbers
            parse.rows = sorted(numbers)
            parse.monthly = monthly
            if parse.statuses and not (set(parse.statuses) & set(numbers)):
                # The chosen rents do not line up with the statuses, so the
                # inclusion filter would silently drop every row.
                parse.status_column, parse.statuses = None, {}
                parse.notes.append("status column ignored: rows do not align with the rent range")

    if parse.rows:
        first_row = min(parse.rows)
        derived = sum(1 for r in parse.rows if wb.formula(sheet, f"{parse.rent_column}{r}"))
        parse.rent_column_is_derived = derived > len(parse.rows) / 2
        header = _header_text(ctx, sheet, parse.rent_column, first_row)
        if _ANNUAL_HEADER.search(header) and not _MONTHLY_HEADER.search(header):
            parse.monthly = False
        elif _MONTHLY_HEADER.search(header):
            parse.monthly = True
        parse.notes.append(f"rent column {parse.rent_column} header: {header[:70]!r}")

        # Note any adjustment columns; their presence means the model's GPR is
        # not a plain annualisation and the difference is for a human to judge.
        index = column_index_from_string(parse.rent_column)
        for offset in range(-12, 13):
            col_index = index + offset
            if col_index < 1:
                continue
            column = get_column_letter(col_index)
            head = _header_text(ctx, sheet, column, first_row)
            if _ADJUSTMENT_HEADER.search(head):
                parse.has_adjustments = True
                parse.adjustment_headers.append(f"{column}: {head[:44]}")
            elif _INCLUDE_HEADER.search(head):
                parse.has_adjustments = True
                parse.adjustment_headers.append(f"{column}: {head[:44]} (analyst include flag)")
    else:
        parse.notes.append("no numeric tenant rent range reachable from the GPR formula")

    return parse


def check_gpr_recompute(ctx: LoanContext) -> list[Finding]:
    """Rebuild annualized in-place rent from the rent roll and compare."""
    wb, tab = ctx.wb, ctx.tab
    coord = tab.cell(Line.GPR)
    if coord is None:
        return []
    model_gpr = wb.number(tab.sheet, coord)
    parse = parse_rent_roll(ctx)
    ctx.facts["rent_roll_parse"] = parse

    if not parse.confident or model_gpr is None:
        return [
            Finding(
                "CHK_GPR_RECOMPUTE",
                Severity.HIGH,
                Status.MANUAL_REVIEW,
                "Gross potential rent could not be independently rebuilt from the rent roll, so "
                "no variance is reported. Rebuild by hand: "
                + ("; ".join(parse.notes) if parse.notes else "rent roll not readable."),
                sheet=tab.sheet,
                cell=coord,
                evidence=f"rent-roll tab: {parse.sheet!r}",
                on_dy_path=True,
            )
        ]

    recomputed = parse.annualized()
    included = len(parse.included_rows())
    basis = "monthly x 12" if parse.monthly else "annual"
    detail = (
        f"{parse.sheet}!{parse.rent_column} over {included} included tenant row(s), {basis}; "
        f"recomputed {recomputed:,.2f} vs model {model_gpr:,.2f}"
        + (f"; status column {parse.status_column}" if parse.status_column else "")
    )
    difference = recomputed - model_gpr
    relative = abs(difference) / max(abs(model_gpr), 1.0)

    if relative <= REVENUE_TOLERANCE:
        # Be explicit about how strong the tie actually is. Where the summed
        # column is computed by the model itself, this confirms the rent roll
        # foots to the OSAR but does not re-derive the per-tenant adjustments
        # inside that column.
        caveat = ""
        if parse.rent_column_is_derived:
            caveat = (
                f" Note: column {parse.rent_column} is calculated by the model rather than raw "
                f"contractual rent, so this confirms the rent roll foots to the OSAR but does "
                f"not independently verify the per-tenant adjustments inside it"
                + (f" ({', '.join(parse.adjustment_headers[:2])})" if parse.adjustment_headers else "")
                + "."
            )
        return [
            Finding(
                "CHK_GPR_RECOMPUTE",
                Severity.HIGH,
                Status.PASS,
                f"Independently rebuilt gross potential rent ties to the model within "
                f"{REVENUE_TOLERANCE:.1%} ({recomputed:,.0f} vs {model_gpr:,.0f})." + caveat,
                sheet=tab.sheet,
                cell=coord,
                evidence=detail,
                on_dy_path=True,
            )
        ]

    if parse.has_adjustments:
        return [
            Finding(
                "CHK_GPR_RECOMPUTE",
                Severity.HIGH,
                Status.MANUAL_REVIEW,
                f"A plain annualisation of in-place rent gives {recomputed:,.0f} against the "
                f"model's {model_gpr:,.0f} ({difference:+,.0f}, {relative:.2%}). The rent roll "
                f"carries step-up, free-rent or include-flag adjustments, which the loan "
                f"definition permits and this tool does not verify - review the difference by "
                f"hand rather than treating it as an error.",
                sheet=tab.sheet,
                cell=coord,
                evidence=f"{detail}; adjustment columns: {parse.adjustment_headers}",
                on_dy_path=True,
            )
        ]

    return [
        Finding(
            "CHK_GPR_RECOMPUTE",
            Severity.HIGH,
            Status.FLAG,
            f"Independently rebuilt gross potential rent differs from the model by "
            f"{difference:+,.0f} ({relative:.2%}), beyond the {REVENUE_TOLERANCE:.1%} tolerance "
            f"for an individual revenue line, and the rent roll shows no step-up or free-rent "
            f"adjustment that would explain it.",
            sheet=tab.sheet,
            cell=coord,
            evidence=detail,
            on_dy_path=True,
        )
    ]


def check_vacancy_recompute(ctx: LoanContext) -> list[Finding]:
    """Rebuild the vacancy line from occupancy and the agreement's floor."""
    wb, tab, params = ctx.wb, ctx.tab, ctx.params
    coord = tab.cell(Line.VACANCY)
    gpr_coord = tab.cell(Line.GPR)
    if coord is None or gpr_coord is None or params is None or not params.vacancy_floor.found:
        return [
            Finding(
                "CHK_VACANCY_RECOMPUTE",
                Severity.HIGH,
                Status.UNVERIFIABLE,
                "Vacancy could not be recomputed: the vacancy line, GPR line or the "
                "agreement's floor is missing.",
                sheet=tab.sheet,
                cell=coord,
            )
        ]

    floor = float(params.vacancy_floor.value)
    threshold = 1.0 - floor
    occupancy = None
    occ_row = tab.rows.get(Line.OCCUPANCY)
    if occ_row:
        for column in (tab.dy_column, "H"):
            occupancy = wb.number(tab.sheet, f"{column}{occ_row}")
            if occupancy is not None:
                break
    gpr = wb.number(tab.sheet, gpr_coord)
    if occupancy is None or gpr is None:
        return [
            Finding(
                "CHK_VACANCY_RECOMPUTE",
                Severity.HIGH,
                Status.MANUAL_REVIEW,
                "Occupancy or GPR could not be read, so the vacancy line was not recomputed.",
                sheet=tab.sheet,
                cell=coord,
                on_dy_path=True,
            )
        ]

    # In-place GPR already reflects actual vacancy, so only the shortfall to the
    # floor is deducted, and only when occupancy is above the threshold.
    if occupancy > threshold and occupancy > 0:
        expected = -(gpr / occupancy) * (floor - (1 - occupancy))
    else:
        expected = 0.0

    model_value = wb.number(tab.sheet, coord)
    if model_value is None:
        model_value = 0.0  # a blank vacancy line is no deduction
    difference = expected - model_value
    scale = max(abs(gpr), 1.0)
    evidence = (
        f"occupancy {occupancy:.4%}, floor {floor:.2%}, GPR {gpr:,.2f} -> expected "
        f"{expected:,.2f}; model {tab.sheet}!{coord} = {model_value:,.2f}"
    )

    if abs(difference) / scale <= REVENUE_TOLERANCE:
        return [
            Finding(
                "CHK_VACANCY_RECOMPUTE",
                Severity.HIGH,
                Status.PASS,
                (
                    f"Vacancy line ties to an independent recompute ({expected:,.0f})."
                    if expected
                    else f"Vacancy is correctly 0: actual vacancy {1 - occupancy:.2%} already "
                    f"exceeds the {floor:.2%} floor and is embedded in in-place rent."
                ),
                sheet=tab.sheet,
                cell=coord,
                evidence=evidence,
                on_dy_path=True,
            )
        ]

    return [
        Finding(
            "CHK_VACANCY_RECOMPUTE",
            Severity.HIGH,
            Status.FLAG,
            f"Vacancy line differs from an independent recompute by {difference:+,.0f}: expected "
            f"{expected:,.0f} at {occupancy:.2%} occupancy against the {floor:.2%} floor, model "
            f"shows {model_value:,.0f}.",
            sheet=tab.sheet,
            cell=coord,
            evidence=evidence,
            on_dy_path=True,
        )
    ]


RECOMPUTE_CHECKS = (check_gpr_recompute, check_vacancy_recompute)


def run_recompute(ctx: LoanContext) -> list[Finding]:
    findings: list[Finding] = []
    for check in RECOMPUTE_CHECKS:
        findings.extend(check(ctx))
    return findings

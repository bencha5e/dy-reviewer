"""The LLM revenue review: prompt, parsing, failure paths, and the merge.

No test here touches the network. The Anthropic client is stubbed, so what is
under test is our half of the contract - the request we build, the response
shapes we accept, and what reaches the report - not the model's judgment.

The house rule that every check is asserted in both directions applies to the
merge too: switching the review on must remove the deterministic revenue
findings *and* leave every other finding exactly where it was.
"""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from dy_audit import llm
from dy_audit.audit import audit_loan
from dy_audit.discovery import discover
from dy_audit.model import Severity, Status
from dy_audit.report import write_findings_workbook

# --------------------------------------------------------------------------
# Stub client
# --------------------------------------------------------------------------


def _text(payload) -> SimpleNamespace:
    body = payload if isinstance(payload, str) else json.dumps(payload)
    return SimpleNamespace(type="text", text=body)


def _tool_block() -> SimpleNamespace:
    """A server-tool block of the kind code execution puts around the answer."""
    return SimpleNamespace(type="server_tool_use", name="code_execution", input={})


def _message(content, stop_reason="end_turn", **usage) -> SimpleNamespace:
    return SimpleNamespace(
        content=content,
        stop_reason=stop_reason,
        usage=SimpleNamespace(
            input_tokens=usage.get("input_tokens", 100),
            output_tokens=usage.get("output_tokens", 20),
            cache_read_input_tokens=usage.get("cache_read_input_tokens", 0),
        ),
    )


class _Stream:
    def __init__(self, message):
        self._message = message

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def get_final_message(self):
        return self._message


class FakeClient:
    """Records what we sent and replays a canned sequence of responses."""

    def __init__(self, responses, upload_error=None):
        self._responses = list(responses)
        self._upload_error = upload_error
        self.requests: list[dict] = []
        self.uploaded: list[str] = []
        self.deleted: list[str] = []

        outer = self

        class _Files:
            def upload(self, file, **kwargs):
                if outer._upload_error:
                    raise outer._upload_error
                outer.uploaded.append(file[0])
                return SimpleNamespace(id="file_stub_1")

            def delete(self, file_id):
                outer.deleted.append(file_id)

        class _Messages:
            def stream(self, **kwargs):
                outer.requests.append(kwargs)
                if not outer._responses:
                    raise AssertionError("stub ran out of responses")
                return _Stream(outer._responses.pop(0))

        self.beta = SimpleNamespace(files=_Files())
        self.messages = _Messages()


FINDING = {
    "id": "F1",
    "severity": "Blocker",
    "type": "Quantified",
    "revenue_line": "Gross Potential Rent",
    "rule": "R-01",
    "check_id": "CHK_GPR_RECOMPUTE",
    "cells": ["OSAR!I12", "'1Q26 RR'!V7"],
    "clause": "the annualized gross residential rental income",
    "finding": "GPR overstated by 572,580.",
    "impact_annualized": 572580.0,
    "evidence": "=SUM('1Q26 RR'!V7:V300)",
    "recommendation": "Remove the applicant rows from the GPR sum.",
}

TERM_SHEET = {
    "delinquency_threshold_days": 45,
    "notice_to_vacate": "no time limit",
    "vacancy_floor_pct": 0.03,
    "vacancy_floor_base": "base rent",
    "occupancy_cap_pct": None,
    "new_lease_window_days": 90,
    "concessions_basis": "T6 annualized",
    "other_income_basis": "T12",
}


def payload(*findings, **overrides) -> dict:
    """A full review payload. Most tests care about one field of it."""
    body = {
        "loan": "Stub Loan",
        "test_date": "2026-03-31",
        "osar_tab": "OSAR",
        "term_sheet": dict(TERM_SHEET),
        "rebuild": [
            {
                "line": "Gross Potential Rent",
                "osar_cell": "I12",
                "reported": 5_000_000.0,
                "rebuilt": 4_427_420.0,
                "variance": 572_580.0,
                "flag": "FLAG",
                "derivation": "Summed column V over rows 7-300, excluding applicants.",
            }
        ],
        "findings": list(findings) or [FINDING],
        "reviewer_note": "The GPR rebuild did not tie. Applicant rows are counted.",
    }
    body.update(overrides)
    return body


#: A stub transport that returns a full payload, for the merge tests.
def _stub_call(*a, **k):
    return payload(), {"input_tokens": 1, "output_tokens": 1}


@pytest.fixture
def rules(tmp_path, monkeypatch) -> Path:
    path = tmp_path / "rules.md"
    path.write_text("# Revenue rules\nTrace GPR to tenant rows.\n", encoding="utf-8")
    monkeypatch.setenv("DY_REVENUE_RULES", str(path))
    return path


# --------------------------------------------------------------------------
# Parsing
# --------------------------------------------------------------------------


def test_payload_is_found_behind_server_tool_blocks():
    # A code-execution turn interleaves tool blocks with text, so the answer is
    # not at content[0]. Scanning from the end must still find it.
    content = [_text("Reading the workbook now."), _tool_block(), _text(payload())]
    assert llm.extract_payload(content)["findings"][0]["check_id"] == "CHK_GPR_RECOMPUTE"


def test_payload_is_found_inside_a_fenced_block():
    # The fallback for a turn that ignores the schema and writes prose + JSON.
    fenced = "Here is what I found:\n```json\n" + json.dumps(payload()) + "\n```"
    assert llm.extract_payload([_text(fenced)])["findings"][0]["cells"][0] == "OSAR!I12"


def test_unparseable_response_raises():
    with pytest.raises(ValueError, match="no parseable findings"):
        llm.extract_payload([_text("I could not open the file.")])


def test_prose_that_merely_contains_a_brace_is_not_mistaken_for_the_payload():
    with pytest.raises(ValueError):
        llm.extract_payload([_text("{not json at all")])


def test_findings_round_trip_into_the_dataclass():
    findings = llm.to_findings(payload())
    assert len(findings) == 1
    finding = findings[0]
    assert finding.check_id == "CHK_GPR_RECOMPUTE"
    assert finding.severity is Severity.BLOCKER
    assert finding.status is Status.FLAG
    # The first cell fills the Sheet and Cell columns...
    assert finding.cell == "I12"
    assert finding.ref == "OSAR!I12"
    # ...the rule number leads the message, and the impact is spelled out...
    assert finding.message.startswith("[R-01] ")
    assert "572,580.00" in finding.message
    # ...and the fields with no column of their own land in evidence.
    assert "Clause: the annualized" in finding.evidence
    assert "Also: '1Q26 RR'!V7" in finding.evidence
    assert "Recommend: Remove the applicant rows" in finding.evidence


def test_a_sheet_qualified_cell_is_split_back_apart():
    # The Findings sheet has separate Sheet and Cell columns; "OSAR!I12" in the
    # cell field would render wrong and break the ref property.
    item = FINDING | {"cells": ["'1Q26 RR'!V7"]}
    finding = llm.to_findings(payload(item))[0]
    assert finding.sheet == "1Q26 RR"
    assert finding.cell == "V7"


def test_a_bare_cell_with_no_sheet_prefix_is_left_alone():
    finding = llm.to_findings(payload(FINDING | {"cells": ["I12"]}))[0]
    assert finding.sheet is None and finding.cell == "I12"


@pytest.mark.parametrize(
    "kind, status",
    [
        ("Verified", Status.PASS),
        ("Memo", Status.PASS),
        ("Unquantified", Status.UNVERIFIABLE),
        ("Control", Status.FLAG),
        ("Methodology", Status.FLAG),
        ("Latent", Status.FLAG),
        ("Quantified", Status.FLAG),
    ],
)
def test_each_finding_type_maps_to_a_status(kind, status):
    assert llm.to_findings(payload(FINDING | {"type": kind}))[0].status is status


def test_a_memo_passes_but_stays_visible_at_info():
    # A permitted inclusion is not a defect, so it passes - but the reviewer is
    # meant to see it, and INFO is counted in the severity totals regardless of
    # status. Passing at HIGH would read as an alarm; passing silently would not
    # read at all.
    finding = llm.to_findings(payload(FINDING | {"type": "Memo", "severity": "High"}))[0]
    assert finding.status is Status.PASS
    assert finding.severity is Severity.INFO


def test_an_unquantified_finding_never_reads_as_a_pass():
    # The house rule: a required test with no supporting data behind it is the
    # exact case that must not look clean.
    item = FINDING | {"type": "Unquantified", "impact_annualized": None}
    finding = llm.to_findings(payload(item))[0]
    assert finding.status is Status.UNVERIFIABLE
    assert "Impact:" not in finding.message


def test_unknown_severity_or_type_falls_back_rather_than_crashing():
    item = FINDING | {"severity": "CATASTROPHIC", "type": "MAYBE"}
    finding = llm.to_findings(payload(item))[0]
    assert finding.severity is Severity.HIGH
    # An unreadable type must never land on PASS.
    assert finding.status is Status.MANUAL_REVIEW


def test_findings_without_a_message_are_dropped():
    assert llm.to_findings(payload(FINDING | {"finding": "   "})) == []


# --------------------------------------------------------------------------
# The call
# --------------------------------------------------------------------------


def test_the_request_carries_the_workbook_the_tool_and_the_cached_rules(rules, tmp_path):
    xlsx = tmp_path / "Loan.xlsx"
    xlsx.write_bytes(b"stub")
    client = FakeClient([_message([_text(payload())])])

    parsed, usage = llm._call_anthropic(
        xlsx, "RULES TEXT", "BRIEF TEXT", model="claude-opus-5", effort="xhigh", client=client
    )

    assert parsed["findings"][0]["check_id"] == "CHK_GPR_RECOMPUTE"
    assert usage["input_tokens"] == 100

    request = client.requests[0]
    assert request["model"] == "claude-opus-5"
    assert request["tools"] == [{"type": "code_execution_20260521", "name": "code_execution"}]
    assert request["output_config"]["effort"] == "xhigh"
    assert request["output_config"]["format"]["schema"] is llm.REVENUE_FINDINGS_SCHEMA
    # The rules are identical across loans, so they must be a cached prefix.
    assert request["system"][0]["cache_control"] == {"type": "ephemeral"}
    assert request["system"][0]["text"] == "RULES TEXT"
    assert "files-api-2025-04-14" in request["extra_headers"]["anthropic-beta"]

    blocks = request["messages"][0]["content"]
    assert blocks[0]["text"] == "BRIEF TEXT"
    # The workbook goes in as a container upload; a document block would reject
    # an .xlsx outright.
    assert blocks[1] == {"type": "container_upload", "file_id": "file_stub_1"}
    assert client.uploaded == ["Loan.xlsx"]


def test_the_uploaded_file_is_always_deleted(rules, tmp_path):
    xlsx = tmp_path / "Loan.xlsx"
    xlsx.write_bytes(b"stub")

    client = FakeClient([_message([_text(payload())])])
    llm._call_anthropic(xlsx, "r", "b", model="m", effort="high", client=client)
    assert client.deleted == ["file_stub_1"]

    # And on the failure path too, or a long run leaks a file per loan.
    client = FakeClient([_message([_text("not json")])])
    with pytest.raises(ValueError):
        llm._call_anthropic(xlsx, "r", "b", model="m", effort="high", client=client)
    assert client.deleted == ["file_stub_1"]


def test_a_paused_turn_is_resumed(rules, tmp_path):
    # The server-side code-execution loop pauses every 10 iterations. Treating
    # the pause as the end silently truncates the review.
    xlsx = tmp_path / "Loan.xlsx"
    xlsx.write_bytes(b"stub")
    paused = _message([_text("still working"), _tool_block()], stop_reason="pause_turn")
    done = _message([_text(payload())])
    client = FakeClient([paused, done])

    parsed, usage = llm._call_anthropic(
        xlsx, "r", "b", model="m", effort="high", client=client
    )

    assert parsed["findings"][0]["check_id"] == "CHK_GPR_RECOMPUTE"
    assert len(client.requests) == 2, "the paused turn must be resumed"
    # The resume hands the paused assistant turn back so the server continues it.
    resumed = client.requests[1]["messages"]
    assert resumed[-1]["role"] == "assistant"
    assert resumed[-1]["content"] == paused.content
    # Usage accumulates across the resumes, not just the final leg.
    assert usage["input_tokens"] == 200


def test_endless_pausing_gives_up_rather_than_looping(rules, tmp_path):
    xlsx = tmp_path / "Loan.xlsx"
    xlsx.write_bytes(b"stub")
    paused = _message([_text("working")], stop_reason="pause_turn")
    client = FakeClient([paused] * (llm.MAX_CONTINUATIONS + 1))
    with pytest.raises(RuntimeError, match="resumes"):
        llm._call_anthropic(xlsx, "r", "b", model="m", effort="high", client=client)


# --------------------------------------------------------------------------
# Failure is visible, never silent
# --------------------------------------------------------------------------


def _ctx_stub(tmp_path):
    """The smallest context `build_brief` can render, so these tests isolate
    the failure being exercised rather than tripping over a missing attribute."""
    return SimpleNamespace(
        loan_name="Stub Loan",
        files=SimpleNamespace(
            xlsx_path=tmp_path / "Loan.xlsx", defs_path=tmp_path / "absent_defs.md"
        ),
        tab=SimpleNamespace(
            sheet="OSAR",
            dy_column="I",
            period_end=None,
            rows={},
            has=lambda line: False,
            cell=lambda line: None,
        ),
        wb=SimpleNamespace(number=lambda *a: None, formula=lambda *a: None),
        params=None,
        facts={},
    )


def test_a_failed_review_reports_unverifiable_rather_than_passing(rules, tmp_path):
    (tmp_path / "Loan.xlsx").write_bytes(b"stub")
    client = FakeClient([], upload_error=RuntimeError("upload refused"))

    findings = llm.review_revenue_with_llm(_ctx_stub(tmp_path), client=client)

    assert len(findings) == 1
    finding = findings[0]
    assert finding.check_id == llm.LLM_CHECK_ID
    # The house rule: a check the tool could not complete never reads as a pass.
    assert finding.status is Status.UNVERIFIABLE
    assert finding.severity is Severity.HIGH
    assert "upload refused" in finding.evidence


def test_missing_rules_are_reported_as_unverifiable(tmp_path, monkeypatch):
    monkeypatch.setenv("DY_REVENUE_RULES", str(tmp_path / "absent.md"))
    findings = llm.review_revenue_with_llm(_ctx_stub(tmp_path), client=FakeClient([]))
    assert findings[0].status is Status.UNVERIFIABLE
    assert "absent.md" in findings[0].evidence


def test_an_empty_review_is_unverifiable_not_a_clean_bill(rules, tmp_path):
    (tmp_path / "Loan.xlsx").write_bytes(b"stub")
    client = FakeClient([_message([_text(payload(findings=[]))])])
    findings = llm.review_revenue_with_llm(_ctx_stub(tmp_path), client=client)
    assert findings[0].status is Status.UNVERIFIABLE


def test_an_unknown_provider_is_unverifiable(tmp_path):
    findings = llm.review_revenue_with_llm(_ctx_stub(tmp_path), provider="nope")
    assert findings[0].status is Status.UNVERIFIABLE


def test_a_failed_review_is_counted_not_quietly_dropped(strada, tmp_path, monkeypatch):
    # The whole point of the UNVERIFIABLE finding is that the reviewer sees the
    # gap. If it does not reach the headline counts, a loan whose revenue review
    # failed summarises exactly like a clean one.
    monkeypatch.setenv("DY_REVENUE_RULES", str(tmp_path / "absent.md"))
    result = audit_loan(strada, use_llm=True)

    unreviewed = [f for f in result.findings if f.check_id == llm.LLM_CHECK_ID]
    assert len(unreviewed) == 1
    assert unreviewed[0].status is Status.UNVERIFIABLE

    counts = result.count_by_severity()
    assert counts.get(Severity.HIGH, 0) >= 1, "the failed review must be counted"

    # The other direction: a review that ran raises no such finding, and the
    # count drops back.
    monkeypatch.delenv("DY_REVENUE_RULES")
    rules_file = tmp_path / "rules.md"
    rules_file.write_text("# rules", encoding="utf-8")
    monkeypatch.setenv("DY_REVENUE_RULES", str(rules_file))
    monkeypatch.setattr(
        llm, "_call_anthropic",
        _stub_call,
    )
    ok = audit_loan(strada, use_llm=True)
    assert not [f for f in ok.findings if f.check_id == llm.LLM_CHECK_ID]


# --------------------------------------------------------------------------
# The merge
# --------------------------------------------------------------------------


@pytest.fixture
def strada(input_copy):
    pairs, _ = discover(input_copy)
    return next(p for p in pairs if "Strada" in p.loan_name)


def test_deterministic_revenue_findings_give_way_to_the_review(strada, rules, monkeypatch):
    baseline = audit_loan(strada, use_llm=False)
    baseline_ids = {f.check_id for f in baseline.findings}
    assert baseline_ids & llm.REVENUE_CHECK_IDS, "fixture must exercise the revenue checks"

    monkeypatch.setattr(
        llm, "_call_anthropic",
        _stub_call,
    )

    hybrid = audit_loan(strada, use_llm=True)
    hybrid_ids = [f.check_id for f in hybrid.findings]

    # Both directions. The revenue IDs the deterministic suite raised are gone...
    surviving = set(hybrid_ids) & (llm.REVENUE_CHECK_IDS - {"CHK_GPR_RECOMPUTE"})
    assert not surviving, f"deterministic revenue findings leaked through: {surviving}"
    # ...the review's own finding is present...
    assert hybrid_ids.count("CHK_GPR_RECOMPUTE") == 1
    # ...and nothing outside revenue moved.
    assert {f.check_id for f in baseline.findings if f.check_id not in llm.REVENUE_CHECK_IDS} == {
        f.check_id for f in hybrid.findings if f.check_id not in llm.REVENUE_CHECK_IDS
    }


def test_the_rebuild_tab_survives_the_handover(strada, rules, monkeypatch):
    # The revenue checks stop reporting but must keep producing data: the Rent
    # Roll Rebuild sheet is rendered from facts they populate.
    monkeypatch.setattr(
        llm, "_call_anthropic",
        _stub_call,
    )
    result = audit_loan(strada, use_llm=True)
    assert result.facts["rebuilds"], "the rent-roll rebuild must still run"
    assert result.facts["source_tabs"].get("rent_roll")


def test_review_findings_render_into_the_existing_workbook(strada, rules, tmp_path, monkeypatch):
    monkeypatch.setattr(
        llm, "_call_anthropic",
        _stub_call,
    )
    result = audit_loan(strada, use_llm=True)
    path = write_findings_workbook(result, tmp_path / "out.xlsx")

    book = openpyxl.load_workbook(path)
    assert book.sheetnames == [
        "Summary",
        "Findings",
        "Rent Roll Rebuild",
        "Revenue Review",
        "Loan Parameters",
    ]
    sheet = book["Findings"]
    rows = [
        [cell.value for cell in row]
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row)
    ]
    match = [r for r in rows if r[2] == "CHK_GPR_RECOMPUTE"]
    assert len(match) == 1
    row = match[0]
    assert row[0] == "BLOCKER" and row[1] == "FLAG"
    assert row[3] == "OSAR" and row[4] == "I12"
    assert row[5] == "yes"
    assert "572,580" in row[6]
    # A blocker from the review sorts to the top alongside deterministic ones.
    assert rows[0][0] == "BLOCKER"


def test_the_revenue_review_sheet_carries_what_the_findings_sheet_cannot(
    strada, rules, tmp_path, monkeypatch
):
    # The term sheet, the rebuild table and the reviewer note have no column on
    # the Findings sheet. Dropping them would throw away most of what the review
    # returns, so they must land here in full.
    monkeypatch.setattr(llm, "_call_anthropic", _stub_call)
    result = audit_loan(strada, use_llm=True)
    path = write_findings_workbook(result, tmp_path / "out.xlsx")

    sheet = openpyxl.load_workbook(path)["Revenue Review"]
    text = "\n".join(
        str(cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "Applicant rows are counted" in text          # the reviewer note
    assert "Delinquency threshold (days)" in text and "45" in text  # the term sheet
    assert "Summed column V over rows 7-300" in text     # the rebuild derivation
    assert "R-01" in text                                # the rule number
    assert "Remove the applicant rows" in text           # the recommendation
    # A null term-sheet row must say so rather than render blank, which would be
    # indistinguishable from a row the review forgot.
    assert "not stated in the agreement" in text


def test_a_null_term_sheet_renders_without_crashing(strada, rules, tmp_path, monkeypatch):
    # Where an agreement states none of these parameters, every row is null.
    # That is a real answer and must render, not raise.
    blank = {key: None for key in TERM_SHEET}
    monkeypatch.setattr(
        llm,
        "_call_anthropic",
        lambda *a, **k: (payload(term_sheet=blank), {"input_tokens": 1, "output_tokens": 1}),
    )
    result = audit_loan(strada, use_llm=True)
    path = write_findings_workbook(result, tmp_path / "out.xlsx")
    sheet = openpyxl.load_workbook(path)["Revenue Review"]
    text = "\n".join(
        str(c.value) for row in sheet.iter_rows() for c in row if c.value is not None
    )
    assert text.count("not stated in the agreement") == len(TERM_SHEET)


def test_the_review_covers_every_revenue_id_the_suite_can_raise():
    # If a new deterministic revenue check is added without adding its ID here,
    # the loan gets reviewed twice for the same defect. Fail loudly instead.
    allowed = set(llm.REVENUE_FINDINGS_SCHEMA["properties"]["findings"]["items"]["properties"][
        "check_id"
    ]["enum"])
    assert llm.REVENUE_CHECK_IDS <= allowed
    assert "CHK_REVENUE_OTHER" in allowed
    # The model-vs-parser disagreement has no deterministic counterpart, so it is
    # filable but must stay out of the dedupe set or it would suppress nothing
    # and silently vanish.
    assert llm.TERM_SHEET_CHECK_ID in allowed
    assert llm.TERM_SHEET_CHECK_ID not in llm.REVENUE_CHECK_IDS
    # Period and month-count stay with Python: they are arithmetic, not judgment.
    assert "CHK_PERIOD" not in llm.REVENUE_CHECK_IDS
    assert "CHK_MONTH_COUNT" not in llm.REVENUE_CHECK_IDS


# --------------------------------------------------------------------------
# The brief
# --------------------------------------------------------------------------


def test_the_brief_names_the_tabs_figures_and_agreement_terms(strada, rules, monkeypatch):
    captured: dict = {}

    def _capture(xlsx, rules_text, brief, **kwargs):
        captured["brief"] = brief
        return payload(), {"input_tokens": 1, "output_tokens": 1}

    monkeypatch.setattr(llm, "_call_anthropic", _capture)
    audit_loan(strada, use_llm=True)

    brief = captured["brief"]
    assert "Strada" in brief
    assert "OSAR tab selected" in brief
    assert "Rent roll tab" in brief
    assert "Vacancy floor" in brief
    # Python's own rebuild goes in as a second opinion, not as the answer.
    assert "second opinion, not the answer" in brief
    assert "GPR" in brief


def test_the_brief_carries_the_agreement_verbatim(strada, rules, monkeypatch):
    # Phase 0 builds the term sheet from the clause text. Sending only the
    # parser's values leaves it with nothing to read, and the rules that hang
    # off the term sheet cannot run at all.
    captured: dict = {}

    def _capture(xlsx, rules_text, brief, **kwargs):
        captured["brief"] = brief
        return payload(), {"input_tokens": 1, "output_tokens": 1}

    monkeypatch.setattr(llm, "_call_anthropic", _capture)
    audit_loan(strada, use_llm=True)

    brief = captured["brief"]
    assert "## DEFINITIONS" in brief
    agreement = strada.defs_path.read_text(encoding="utf-8").strip()
    assert agreement in brief, "the definitions file must go in whole, not summarised"
    # And the parser's reading is offered beside it as a cross-check, not as law.
    assert "Your reading of the clause wins" in brief
    assert "R-27" in brief


def test_a_missing_definitions_file_is_stated_rather_than_faked(tmp_path):
    brief = llm.build_brief(_ctx_stub(tmp_path))
    assert "no loan agreement definitions file was found" in brief


def test_meta_reads_the_property_count_from_the_overview_block(strada, rules, monkeypatch):
    # The count sits in column E of the overview rows, not in the debt-yield
    # column - reading it from the DY column returns nothing on every loan, and
    # META silently degrades to "not read" without failing anything.
    captured: dict = {}

    def _capture(xlsx, rules_text, brief, **kwargs):
        captured["brief"] = brief
        return payload(), {"input_tokens": 1, "output_tokens": 1}

    monkeypatch.setattr(llm, "_call_anthropic", _capture)
    audit_loan(strada, use_llm=True)

    meta = captured["brief"].split("## META")[1].split("## DEFINITIONS")[0]
    assert "Net rentable units / SF: 495" in meta
    assert "Property type: multifamily (inferred)" in meta
    assert "Test date: 2026-03-31" in meta

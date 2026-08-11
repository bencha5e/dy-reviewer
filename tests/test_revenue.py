"""Acceptance tests for the independent rent-roll rebuild and its companions.

Three workbooks under `tests/fixtures/revenue/` carry the revenue defects this
round was built from:

- **Memorial Hills** - GPR = (RENT + PET RENT + GAS)*12 - T6 concessions: pet
  rent and gas reimbursement are supplementary billing inside the base-rent
  line, and the rent roll is stamped a year stale (As of 03/31/2025).
- **Quincy Hollingsworth** - nine cells in the exported rent column were
  overwritten with `=+P<row>`, pulling LAST quarter's rent onto units the
  export shows producing nothing (+$554,700/yr); the export's own totals no
  longer tie; the OSAR note claims delinquents and known vacates are excluded
  while the formula excludes nobody; and the vacancy formula references a
  blank cell, so it can never deduct.
- **Hialeah (corrected variant)** - the repo fixture's single-row `*0.95`
  haircut on `'1Q26 RR'!V7` removed; the rebuild must tie cleanly here while
  still flagging the repo fixture.

Everything is asserted in both directions: the five main fixtures pin their
rebuild statuses too, so a rule loosened to catch these three cannot silently
start firing on clean loans.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from dy_audit.audit import audit_loan
from dy_audit.discovery import discover
from dy_audit.model import Severity, Status
from dy_audit.osar import Line

REPO_ROOT = Path(__file__).resolve().parent.parent
REVENUE_DIR = REPO_ROOT / "tests" / "fixtures" / "revenue"

FLAG, PASS = Status.FLAG, Status.PASS
MANUAL = Status.MANUAL_REVIEW


@pytest.fixture(scope="module")
def revenue():
    pairs, problems = discover(REVENUE_DIR)
    assert problems == []
    return {p.loan_name: audit_loan(p) for p in pairs}


@pytest.fixture(scope="module")
def main(  # the five main fixtures, for the must-not-fire direction
):
    pairs, problems = discover(REPO_ROOT)
    assert problems == []
    return {p.loan_name: audit_loan(p) for p in pairs}


def _findings(results, loan, check_id):
    return [f for f in results[loan].findings if f.check_id == check_id]


def _one(results, loan, check_id, line=None):
    found = _findings(results, loan, check_id)
    if line is not None:
        found = [f for f in found if f.message.startswith(line)]
    assert len(found) == 1, f"{loan} {check_id}: {[(f.status, f.message[:60]) for f in found]}"
    return found[0]


# -- Memorial Hills ------------------------------------------------------------


def test_memorial_supplementary_income_in_gpr_flags(revenue):
    finding = _one(revenue, "Memorial Hills", "CHK_GPR_SUPPLEMENTARY")
    assert finding.status is FLAG
    assert "PET RENT" in finding.message and "GAS" in finding.message
    assert "54,414" in finding.message
    # The operating statement codes the same charges to Other Income, so the
    # double-count risk is called out with the account it found.
    assert "counted twice" in finding.message


def test_memorial_rebuild_reconciles_to_the_dollar(revenue):
    finding = _one(revenue, "Memorial Hills", "CHK_RENT_REBUILD")
    assert finding.status is FLAG
    rebuild = revenue["Memorial Hills"].facts["rebuilds"][0]
    assert rebuild.columns["rent"] == "R"
    assert len(rebuild.included_rows()) == 265  # occupied family only
    assert rebuild.annualized() == pytest.approx(505_260 * 12)
    # model = rebuilt + pet + gas - concessions, so the residual line is absent.
    items = dict.fromkeys(item for item, _amt in rebuild.reconciliation)
    assert not any("unexplained" in item for item in items)


def test_memorial_stale_rent_roll_flags_the_period(revenue):
    stale = [
        f
        for f in _findings(revenue, "Memorial Hills", "CHK_PERIOD")
        if f.status is FLAG and f.sheet == "Rent Roll"
    ]
    assert stale, "the 03/31/2025 rent roll must flag against the 03/31/2026 quarter"


def test_memorial_vacancy_formula_is_correct_and_passes(revenue):
    # `IF(O16>O17,"",(O17-O16)*-I24)` deducts below the floor and blanks above
    # it - correct. The driver finder must substitute occupancy into I16, not
    # into O16 (= 1-I16, the vacancy rate on the same row).
    assert _one(revenue, "Memorial Hills", "CHK_VACANCY_SIGN").status is PASS


def test_memorial_export_total_ties_on_the_occupied_basis(revenue):
    # The export's Totals row nets applicant and former-resident rent - an
    # exclusion basis, not tampering.
    assert _one(revenue, "Memorial Hills", "CHK_RR_TOTAL_TIE").status is PASS


# -- Quincy Hollingsworth ------------------------------------------------------


def test_quincy_backfilled_rent_cells_flag(revenue):
    flags = [
        f
        for f in _findings(revenue, "Quincy and Hollingsworth", "CHK_RENT_EDITS")
        if f.status is FLAG
    ]
    assert len(flags) == 1
    finding = flags[0]
    assert "9 cell(s)" in finding.message
    assert "554,700" in finding.message
    assert "another period" in finding.evidence  # =+P51 pulls the 4Q25 column


def test_quincy_export_totals_no_longer_tie(revenue):
    flags = [
        f
        for f in _findings(revenue, "Quincy and Hollingsworth", "CHK_RR_TOTAL_TIE")
        if f.status is FLAG
    ]
    assert len(flags) == 1
    assert "edited after export" in flags[0].message
    # The commercial schedule's totals still tie and must stay a PASS.
    passes = [
        f
        for f in _findings(revenue, "Quincy and Hollingsworth", "CHK_RR_TOTAL_TIE")
        if f.status is PASS
    ]
    assert passes and passes[0].sheet == "Rent Roll (COMM)"


def test_quincy_note_claims_exclusions_the_formula_never_applies(revenue):
    rebuild = next(
        rb
        for rb in revenue["Quincy and Hollingsworth"].facts["rebuilds"]
        if rb.line is Line.GPR
    )
    items = [item for item, _amt in rebuild.reconciliation]
    assert any("delinquent tenants are excluded, but these are in" in i for i in items)
    assert any("known vacates are excluded" in i for i in items)


def test_quincy_commercial_base_rent_rebuild_ties(revenue):
    # Base Rent = the retail block only; the parking block below the retail
    # Total row belongs to the Parking line and must not leak in.
    base = [
        f
        for f in _findings(revenue, "Quincy and Hollingsworth", "CHK_RENT_REBUILD")
        if "BASE_RENT" in f.message
    ]
    assert len(base) == 1 and base[0].status is PASS
    comm = next(
        rb
        for rb in revenue["Quincy and Hollingsworth"].facts["rebuilds"]
        if rb.line is Line.BASE_RENT
    )
    assert comm.annualized() == pytest.approx(2_774_713.44)


def test_quincy_inert_vacancy_formula_is_a_blocker(revenue):
    # `IF((1-N16)>N15,...)` tests a blank cell instead of occupancy, so the
    # line returns 0 at every substituted occupancy. Whichever branch words
    # it, it must be a blocker-severity flag showing the zero-everywhere test.
    finding = _one(revenue, "Quincy and Hollingsworth", "CHK_VACANCY_SIGN")
    assert finding.status is FLAG
    assert finding.severity is Severity.BLOCKER
    assert "100.00% -> 0.00" in finding.evidence or "does not depend on occupancy" in finding.message


def test_quincy_gpr_recompute_handles_the_two_block_sum(revenue):
    # SUM(H401:H475,H9:H388) is one tenant column split in two; the recompute
    # must union the slices instead of scoring them separately.
    assert _one(revenue, "Quincy and Hollingsworth", "CHK_GPR_RECOMPUTE").status is PASS


# -- Hialeah corrected variant -------------------------------------------------


def test_corrected_hialeah_rebuild_ties(revenue):
    finding = _one(revenue, "Hialeah Infill Industrial Park", "CHK_RENT_REBUILD")
    assert finding.status is PASS
    rebuild = revenue["Hialeah Infill Industrial Park"].facts["rebuilds"][0]
    assert rebuild.annualized() == pytest.approx(1_436_451.96)


def test_repo_hialeah_still_flags_the_single_row_haircut(main):
    finding = _one(main, "Hialeah", "CHK_RENT_REBUILD")
    assert finding.status is FLAG
    assert "V7" in finding.message and "(Q7*12)*0.95" in finding.message
    assert "-22,525" in finding.message


# -- the five main fixtures must keep their rebuild statuses -------------------


@pytest.mark.parametrize(
    "loan, rebuild_status, edits_status",
    [
        ("Strada", PASS, PASS),
        ("Campus at Villa La Jolla", PASS, MANUAL),
        ("Ares55thAve", PASS, PASS),
        ("Lydian", FLAG, PASS),
    ],
)
def test_main_fixture_rebuild_matrix(main, loan, rebuild_status, edits_status):
    assert _one(main, loan, "CHK_RENT_REBUILD").status is rebuild_status
    assert _one(main, loan, "CHK_RENT_EDITS").status is edits_status


def test_campus_rebuild_credits_the_documented_adjustments(main):
    finding = _one(main, "Campus at Villa La Jolla", "CHK_RENT_REBUILD")
    assert "documented per-tenant adjustments" in finding.message


def test_lydian_rebuild_and_recompute_agree(main):
    rebuild = main["Lydian"].facts["rebuilds"][0]
    assert rebuild.annualized() == pytest.approx(4_081_248.0)
    items = [item for item, _amt in rebuild.reconciliation]
    assert any("Applicant" in i for i in items)
    assert any("Pending renewal" in i for i in items)


def test_every_findings_workbook_gets_a_rebuild_tab(revenue, tmp_path):
    import datetime as dt

    from openpyxl import load_workbook

    from dy_audit.report import write_findings_workbook

    for name, result in revenue.items():
        path = write_findings_workbook(result, tmp_path / f"{name}.xlsx", dt.date(2026, 8, 11))
        wb = load_workbook(path)
        assert "Rent Roll Rebuild" in wb.sheetnames
        sheet = wb["Rent Roll Rebuild"]
        assert sheet.max_row > 5, f"{name}: rebuild tab is empty"

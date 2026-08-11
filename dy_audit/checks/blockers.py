"""BLOCKER checks - spec section 4.

These five decide whether a reported debt yield can be trusted. Three of the
four Q1 2026 defects they catch produce a *correct* DY this quarter and a wrong
one the moment an input moves, which is why they are checked structurally rather
than by comparing this quarter's numbers.
"""

from __future__ import annotations

import re

from .. import formula as F
from ..context import LoanContext
from ..model import Finding, Severity, Status
from ..osar import Line, normalize_label

#: Occupancy scenarios for the vacancy sign test (spec section 5, rule 4).
#: 0.95 is evaluated and reported but never flagged: it is the exact boundary
#: where "greater of actual or 5%" is satisfied by either reading, and the four
#: models legitimately disagree there.
OCCUPANCY_SCENARIOS = (1.00, 0.97, 0.95, 0.92)
_BOUNDARY = 0.95

#: A vacancy line this close to zero counts as zero (floating point leaves -0.0).
_ZERO_TOL = 0.01

#: An invoice term smaller than this share of the T12 term can never win the
#: MAX, which means the MAX is decorative. Ares's broken term is ~7e-13 of T12,
#: so the threshold is nowhere near any legitimate invoice.
_DEGENERATE_RATIO = 0.01

#: Sheets that carry the invoice side of each MAX.
_INVOICE_SHEET = {
    Line.TAX: re.compile(r"tax", re.IGNORECASE),
    Line.INSURANCE: re.compile(r"insur", re.IGNORECASE),
}

#: Labels naming a covenant level rather than a computed debt yield.
_THRESHOLD_WORDS = re.compile(
    r"\b(event|threshold|covenant|minimum|min|required|require|hurdle|trigger|target|test level)\b",
    re.IGNORECASE,
)

_DY_LABEL = re.compile(r"(^|\b)(debt yield|dy)\b", re.IGNORECASE)

_ROW_RE = re.compile(r"^([A-Z]+)(\d+)$")


def _row_of(coord: str) -> int | None:
    m = _ROW_RE.match(coord.replace("$", "").upper())
    return int(m.group(2)) if m else None


# ---------------------------------------------------------------------------
# CHK_TAX_MAX / CHK_INS_MAX
# ---------------------------------------------------------------------------


def _check_max_line(ctx: LoanContext, line: Line, check_id: str, label: str) -> list[Finding]:
    """Shared engine for `Taxes = MAX(T12, invoice)` and the insurance twin.

    Universal hard rule (spec section 2.2): a model referencing only one side is
    a blocker on every loan, not a per-loan option.
    """
    wb, tab = ctx.wb, ctx.tab
    coord = tab.cell(line)
    if coord is None:
        return [
            Finding(
                check_id,
                Severity.BLOCKER,
                Status.UNVERIFIABLE,
                f"No {label} line found on the OSAR tab; MAX(T12, invoice) could not be verified.",
                sheet=tab.sheet,
            )
        ]

    res = F.resolve_defining_formula(wb, tab.sheet, coord)
    # Ares routes the OSAR line through a rollup, so cite where the logic lives.
    where = f"{tab.sheet}!{coord}"
    via = f" (via {' -> '.join(res.path)})" if res.hopped else ""

    if res.formula is None:
        return [
            Finding(
                check_id,
                Severity.BLOCKER,
                Status.FLAG,
                f"{label} is a hardcoded value, not MAX(T12 actual, invoice). "
                f"The invoice can never be picked up.",
                sheet=tab.sheet,
                cell=coord,
                evidence=f"{where}{via} = {res.value!r}",
                on_dy_path=True,
            )
        ]

    args = F.extract_call_args(res.formula, "MAX")
    if args is None:
        return [
            Finding(
                check_id,
                Severity.BLOCKER,
                Status.FLAG,
                f"{label} does not apply MAX(T12 actual, invoice) - it references only one side. "
                f"Correct this quarter only if that side happens to be the larger.",
                sheet=tab.sheet,
                cell=coord,
                evidence=f"{res.ref}{via} = {F.normalize(res.formula)}",
                on_dy_path=True,
            )
        ]

    invoice_pattern = _INVOICE_SHEET[line]
    resolver = F.make_resolver(wb)
    invoice_terms: list[tuple[str, float]] = []
    other_terms: list[tuple[str, float]] = []
    unevaluated: list[str] = []

    for arg in args:
        refs = F.iter_refs(arg)
        touches_invoice = any(
            invoice_pattern.search(r.resolved_sheet(res.sheet)) for r in refs
        )
        try:
            value = F.to_number(F.evaluate(arg, res.sheet, resolver))
        except (F.UnsupportedFormula, F.EvalError) as exc:
            unevaluated.append(f"{arg} ({exc})")
            continue
        (invoice_terms if touches_invoice else other_terms).append((arg, value))

    detail = f"{res.ref}{via} = {F.normalize(res.formula)}"

    if not invoice_terms:
        return [
            Finding(
                check_id,
                Severity.BLOCKER,
                Status.FLAG,
                f"{label} takes a MAX but neither side reads the {label.lower()} invoice tab, "
                f"so the actual bill is never compared.",
                sheet=tab.sheet,
                cell=coord,
                evidence=detail,
                on_dy_path=True,
            )
        ]

    if unevaluated and not other_terms:
        return [
            Finding(
                check_id,
                Severity.BLOCKER,
                Status.MANUAL_REVIEW,
                f"{label} MAX found, but a term could not be evaluated; confirm by hand.",
                sheet=tab.sheet,
                cell=coord,
                evidence=f"{detail} | unevaluated: {'; '.join(unevaluated)}",
                on_dy_path=True,
            )
        ]

    invoice_arg, invoice_value = max(invoice_terms, key=lambda t: abs(t[1]))
    comparison = max((abs(v) for _, v in other_terms), default=0.0)

    if comparison > 0 and abs(invoice_value) < _DEGENERATE_RATIO * comparison:
        return [
            Finding(
                check_id,
                Severity.BLOCKER,
                Status.FLAG,
                f"{label} applies MAX, but the invoice term resolves to "
                f"{invoice_value:,.6g} against a T12 actual of {comparison:,.2f} - it can never "
                f"win the MAX, so the actual premium is effectively ignored. Check for a "
                f"unit-scaling error (a value displayed in millions divided by its scale cell).",
                sheet=tab.sheet,
                cell=coord,
                evidence=f"{detail} | invoice term `{invoice_arg}` = {invoice_value!r}",
                on_dy_path=True,
            )
        ]

    return [
        Finding(
            check_id,
            Severity.BLOCKER,
            Status.PASS,
            f"{label} correctly applies MAX(T12 actual, invoice); "
            f"invoice term = {invoice_value:,.2f}, T12 term = {comparison:,.2f}.",
            sheet=tab.sheet,
            cell=coord,
            evidence=detail,
            on_dy_path=True,
        )
    ]


def check_tax_max(ctx: LoanContext) -> list[Finding]:
    return _check_max_line(ctx, Line.TAX, "CHK_TAX_MAX", "Real Estate Taxes")


def check_ins_max(ctx: LoanContext) -> list[Finding]:
    return _check_max_line(ctx, Line.INSURANCE, "CHK_INS_MAX", "Property Insurance")


# ---------------------------------------------------------------------------
# CHK_VACANCY_SIGN
# ---------------------------------------------------------------------------


def _find_occupancy_driver(ctx: LoanContext, res: F.Resolution) -> tuple[str | None, dict]:
    """Locate the occupancy cell the vacancy formula actually depends on.

    Ares reads occupancy from column H of the OSAR tab rather than the DY
    column, so the driver is found by matching the occupancy *row* anywhere on
    the OSAR sheet within the formula's dependency closure.
    """
    deps = F.dependencies(ctx.wb, res.sheet, res.coord)
    occ_row = ctx.tab.rows.get(Line.OCCUPANCY)
    if occ_row is None:
        return None, deps
    for key, (sheet, coord) in deps.items():
        if sheet == ctx.tab.sheet and _row_of(coord) == occ_row:
            return key, deps
    return None, deps


def check_vacancy_sign(ctx: LoanContext) -> list[Finding]:
    """Vacancy loss must be a deduction above 95% occupancy and zero below it.

    Evaluated under substituted occupancy because the defect is invisible at the
    current figure: Strada's inverted branch returns 0 at 91.95% occupancy.
    """
    wb, tab = ctx.wb, ctx.tab
    coord = tab.cell(Line.VACANCY)
    if coord is None:
        return [
            Finding(
                "CHK_VACANCY_SIGN",
                Severity.BLOCKER,
                Status.UNVERIFIABLE,
                "No vacancy-loss line found on the OSAR tab.",
                sheet=tab.sheet,
            )
        ]

    res = F.resolve_defining_formula(wb, tab.sheet, coord)
    where = f"{tab.sheet}!{coord}"
    via = f" (via {' -> '.join(res.path)})" if res.hopped else ""

    if res.formula is None:
        value = F.to_number(res.value) if res.value is not None else 0.0
        status = Status.FLAG if value > _ZERO_TOL else Status.MANUAL_REVIEW
        return [
            Finding(
                "CHK_VACANCY_SIGN",
                Severity.BLOCKER,
                status,
                "Vacancy loss is a hardcoded value, so the 5% floor cannot respond to "
                "occupancy. Confirm by hand.",
                sheet=tab.sheet,
                cell=coord,
                evidence=f"{where}{via} = {res.value!r}",
                on_dy_path=True,
            )
        ]

    occ_key, deps = _find_occupancy_driver(ctx, res)
    if occ_key is None:
        return [
            Finding(
                "CHK_VACANCY_SIGN",
                Severity.BLOCKER,
                Status.MANUAL_REVIEW,
                "Could not identify the occupancy cell the vacancy formula depends on, so the "
                "sign could not be tested at synthetic occupancies. Review by hand.",
                sheet=tab.sheet,
                cell=coord,
                evidence=f"{res.ref}{via} = {F.normalize(res.formula)}",
                on_dy_path=True,
            )
        ]
    ctx.facts["occupancy_driver"] = occ_key

    # Recompute only what actually moves with occupancy; everything else keeps
    # the value Excel cached, which keeps lookups off the evaluation path.
    dependents = {occ_key}
    for key, (sheet, cell) in deps.items():
        if occ_key in F.dependencies(wb, sheet, cell):
            dependents.add(key)

    results: dict[float, float] = {}
    for occ in OCCUPANCY_SCENARIOS:
        resolver = F.make_resolver(wb, overrides={occ_key: occ}, recompute=dependents)
        try:
            results[occ] = F.to_number(F.evaluate(res.formula, res.sheet, resolver))
        except (F.UnsupportedFormula, F.EvalError) as exc:
            return [
                Finding(
                    "CHK_VACANCY_SIGN",
                    Severity.BLOCKER,
                    Status.MANUAL_REVIEW,
                    f"Vacancy formula could not be evaluated at {occ:.0%} occupancy "
                    f"({exc}); review the sign by hand.",
                    sheet=tab.sheet,
                    cell=coord,
                    evidence=f"{res.ref}{via} = {F.normalize(res.formula)}",
                    on_dy_path=True,
                )
            ]

    trace = ", ".join(f"{occ:.0%} -> {val:,.2f}" for occ, val in results.items())
    evidence = f"{res.ref}{via} = {F.normalize(res.formula)} | occupancy test: {trace}"
    problems: list[str] = []

    for occ, val in results.items():
        if occ == _BOUNDARY:
            continue  # exact-floor boundary: reported, never flagged
        if occ > _BOUNDARY and val > -_ZERO_TOL:
            problems.append(
                f"at {occ:.0%} occupancy the line is {val:,.2f}, which "
                f"{'raises' if val > _ZERO_TOL else 'does not reduce'} EGI instead of deducting"
            )
        elif occ < _BOUNDARY and abs(val) > _ZERO_TOL:
            problems.append(
                f"at {occ:.0%} occupancy the line is {val:,.2f}, but actual vacancy already "
                f"exceeds the 5% floor so it must be 0"
            )

    if problems:
        return [
            Finding(
                "CHK_VACANCY_SIGN",
                Severity.BLOCKER,
                Status.FLAG,
                "Vacancy-loss formula is wrong outside the current occupancy: "
                + "; ".join(problems)
                + ". The reported DY may be unaffected this quarter and wrong the next.",
                sheet=tab.sheet,
                cell=coord,
                evidence=evidence,
                on_dy_path=True,
            )
        ]

    return [
        Finding(
            "CHK_VACANCY_SIGN",
            Severity.BLOCKER,
            Status.PASS,
            "Vacancy loss deducts above 95% occupancy and is zero below it.",
            sheet=tab.sheet,
            cell=coord,
            evidence=evidence,
            on_dy_path=True,
        )
    ]


# ---------------------------------------------------------------------------
# CHK_DY_BASIS
# ---------------------------------------------------------------------------


def check_dy_basis(ctx: LoanContext) -> list[Finding]:
    """The DY numerator must be NCF (post-reserve), not the pre-reserve NOI line.

    Every loan's definitional NOI already nets a replacement reserve, which these
    models place below the NOI line inside Capital Items - so the model's NCF is
    the definitional NOI. Driving DY off the NOI line overstates it.
    """
    wb, tab = ctx.wb, ctx.tab
    coord = tab.cell(Line.DEBT_YIELD)
    if coord is None:
        return [
            Finding(
                "CHK_DY_BASIS",
                Severity.BLOCKER,
                Status.UNVERIFIABLE,
                "No Debt Yield line found on the OSAR tab.",
                sheet=tab.sheet,
            )
        ]

    body = F.normalize(wb.formula(tab.sheet, coord))
    if not body:
        return [
            Finding(
                "CHK_DY_BASIS",
                Severity.BLOCKER,
                Status.FLAG,
                "Debt Yield is a hardcoded value rather than NCF / UPB.",
                sheet=tab.sheet,
                cell=coord,
                evidence=f"{tab.sheet}!{coord} = {wb.value(tab.sheet, coord)!r}",
                on_dy_path=True,
            )
        ]

    refs = [r for r in F.iter_refs(body) if not r.is_range]
    if len(refs) < 2 or "/" not in body:
        return [
            Finding(
                "CHK_DY_BASIS",
                Severity.BLOCKER,
                Status.MANUAL_REVIEW,
                "Debt Yield formula is not a simple numerator/denominator; confirm the "
                "numerator is NCF by hand.",
                sheet=tab.sheet,
                cell=coord,
                evidence=f"{tab.sheet}!{coord} = {body}",
                on_dy_path=True,
            )
        ]

    numerator, denominator = refs[0], refs[1]
    ctx.facts["upb_ref"] = (denominator.resolved_sheet(tab.sheet), denominator.coord)
    ctx.facts["dy_value"] = wb.number(tab.sheet, coord)

    num_row = _row_of(numerator.coord)
    ncf_row = tab.rows.get(Line.NCF)
    noi_row = tab.rows.get(Line.NOI)
    evidence = f"{tab.sheet}!{coord} = {body} (NCF row {ncf_row}, NOI row {noi_row})"

    if noi_row is not None and num_row == noi_row:
        return [
            Finding(
                "CHK_DY_BASIS",
                Severity.BLOCKER,
                Status.FLAG,
                "Debt Yield is driven off the pre-reserve NOI line instead of NCF, which "
                "omits the replacement reserve and overstates the reported yield.",
                sheet=tab.sheet,
                cell=coord,
                evidence=evidence,
                on_dy_path=True,
            )
        ]

    if ncf_row is None or num_row != ncf_row:
        return [
            Finding(
                "CHK_DY_BASIS",
                Severity.BLOCKER,
                Status.FLAG,
                f"Debt Yield numerator is {numerator}, which is not the NCF line "
                f"(row {ncf_row}). Confirm the correct basis.",
                sheet=tab.sheet,
                cell=coord,
                evidence=evidence,
                on_dy_path=True,
            )
        ]

    return [
        Finding(
            "CHK_DY_BASIS",
            Severity.BLOCKER,
            Status.PASS,
            "Debt Yield = NCF / UPB, taken off the post-reserve NCF line.",
            sheet=tab.sheet,
            cell=coord,
            evidence=evidence,
            on_dy_path=True,
        )
    ]


# ---------------------------------------------------------------------------
# CHK_DY_CONSISTENCY
# ---------------------------------------------------------------------------


def _candidate_dy_cells(ctx: LoanContext) -> list[tuple[str, str, float, str]]:
    """Find computed debt yields on tabs other than the OSAR output tab.

    Three filters keep this from firing on things that are not rival
    calculations: hidden sheets are out of scope entirely, a cell without a
    formula is an input rather than a calculation (Ares parks its 6.50% covenant
    level as a literal), and threshold vocabulary in the label rules out
    covenant levels that happen to be computed.
    """
    from openpyxl.utils import column_index_from_string, get_column_letter

    wb, tab = ctx.wb, ctx.tab
    hits: list[tuple[str, str, float, str]] = []

    for sheet in wb.sheet_names:
        if sheet == tab.sheet or not wb.is_visible(sheet):
            continue
        for coord, _f, value in wb.iter_cells(sheet, max_row=200, max_col=30):
            if not isinstance(value, str):
                continue
            label = normalize_label(value)
            if not _DY_LABEL.search(label) or _THRESHOLD_WORDS.search(label):
                continue
            m = _ROW_RE.match(coord)
            if not m:
                continue
            col_index = column_index_from_string(m.group(1))
            row = m.group(2)
            for step in range(1, 8):
                near = f"{get_column_letter(col_index + step)}{row}"
                near_formula = wb.formula(sheet, near)
                if near_formula is None:
                    continue  # a literal is a threshold or input, not a calc
                number = wb.number(sheet, near)
                if number is None or not (0 < abs(number) < 1.5):
                    continue
                hits.append((sheet, near, number, value.strip()))
                break
    return hits


def check_dy_consistency(ctx: LoanContext, tolerance: float = 1e-4) -> list[Finding]:
    """Off-tab debt yields are checked for tie-out only, never audited.

    The OSAR debt yield is the number of record (spec section 5, rule 8). If
    another tab disagrees, name both values and stop there.
    """
    wb, tab = ctx.wb, ctx.tab
    coord = tab.cell(Line.DEBT_YIELD)
    osar_dy = wb.number(tab.sheet, coord) if coord else None
    if osar_dy is None:
        return [
            Finding(
                "CHK_DY_CONSISTENCY",
                Severity.BLOCKER,
                Status.UNVERIFIABLE,
                "OSAR debt yield has no value, so off-tab calculations cannot be tied out.",
                sheet=tab.sheet,
                cell=coord,
            )
        ]

    findings: list[Finding] = []
    for sheet, cell, value, label in _candidate_dy_cells(ctx):
        if abs(value - osar_dy) <= tolerance:
            continue
        findings.append(
            Finding(
                "CHK_DY_CONSISTENCY",
                Severity.BLOCKER,
                Status.FLAG,
                f"A second debt yield on {sheet!r} disagrees with the OSAR figure of record: "
                f"{sheet}!{cell} = {value:.4%} vs {tab.sheet}!{coord} = {osar_dy:.4%}. "
                f"Decide which is reported and fix the other; the off-tab calculation is not "
                f"audited further.",
                sheet=sheet,
                cell=cell,
                evidence=f"label {label!r}; {sheet}!{cell} = {F.normalize(wb.formula(sheet, cell))}",
                on_dy_path=False,
            )
        )

    if findings:
        return findings

    return [
        Finding(
            "CHK_DY_CONSISTENCY",
            Severity.BLOCKER,
            Status.PASS,
            f"No off-tab debt yield disagrees with the OSAR figure of {osar_dy:.4%}.",
            sheet=tab.sheet,
            cell=coord,
            on_dy_path=True,
        )
    ]


#: Run order for the blocker suite.
BLOCKER_CHECKS = (
    check_tax_max,
    check_ins_max,
    check_vacancy_sign,
    check_dy_basis,
    check_dy_consistency,
)


def run_blockers(ctx: LoanContext) -> list[Finding]:
    findings: list[Finding] = []
    for check in BLOCKER_CHECKS:
        findings.extend(check(ctx))
    return findings
